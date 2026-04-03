###################################################################################################
# Babson College AI Grading Assistant
# VERSION 3 — Prompt Engineering Update
#
# Original capabilities retained:
# ✅ Robust error handling with explicit, readable log messages
# ✅ JSON self-repair system
# ✅ Auto-adds missing keys, scores, justifications
# ✅ Auto-corrects D/F award logic
# ✅ Auto-corrects totals
# ✅ Fallback retry system with clear logging
# ✅ Checkpoint saves every 5 rows
#
# Changes in this version:
# FIX 1  — read_docx_text now extracts table cell content (rubric was silently incomplete)
# FIX 2  — temperature set to 0 for deterministic, reproducible grading
# FIX 3  — fabricated quote injection removed; missing justifications flagged for review
# FIX 4  — Loop_Type (Balancing/Reinforcing) added to JSON schema and prompt
# FIX 5  — Bounded inference: evidence-based justifications replace literal quote requirement
# FIX 6  — Explicit loop-closure reasoning test added for Point D
# FIX 7  — Point E exclusion language added (fisher-feeds-themselves does not earn E)
# FIX 8  — Rubric preamble codifying inference permission and D closure rule
# FIX 9  — Borderline score flagging (Needs_Review column + log warnings)
#
###################################################################################################

import json
import logging
import os
import re
import time
from dataclasses import dataclass
from datetime import datetime
from typing import Any, Dict, List

import pandas as pd
from docx import Document
from dotenv import load_dotenv
from openai import AzureOpenAI
from tenacity import retry, stop_after_attempt, wait_exponential

###################################################################################################
# Logging Setup
###################################################################################################

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s - %(levelname)s - %(message)s",
    handlers=[
        logging.FileHandler("grading_log.txt", mode='a', encoding="utf-8"),
        logging.StreamHandler()
    ]
)
logger = logging.getLogger(__name__)

load_dotenv()
DEFAULT_MODEL = os.getenv("AZURE_OPENAI_DEPLOYMENT", "gpt-4.1-mini")

###################################################################################################
# Helper Classes & Paths
###################################################################################################

@dataclass
class Paths:
    input_xlsx: str
    rubric_docx: str
    examples_xlsx: str
    output_xlsx: str
    checkpoint_xlsx: str

###################################################################################################
# File Reading Utilities
###################################################################################################

# FIX 1: Original function only read doc.paragraphs, which skips all table cell content.
# The rubric is almost entirely structured as Word tables (criteria, examples, comments).
# This version collects paragraphs first, then iterates all tables to capture cell text.
def read_docx_text(path: str) -> str:
    """Load rubric docx text into a single normalized string, including all table content."""
    doc = Document(path)
    parts = []

    # Body paragraphs (introductory text, headers, footnotes outside tables)
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            parts.append(text)

    # Table content — this is where the A-F rubric criteria actually live
    for table in doc.tables:
        for row in table.rows:
            row_parts = []
            for cell in row.cells:
                cell_text = cell.text.strip()
                if cell_text:
                    row_parts.append(cell_text)
            if row_parts:
                parts.append(" | ".join(row_parts))
        parts.append("")  # blank line between tables for readability

    extracted = "\n".join(parts)
    logger.info(f"Rubric extracted: {len(extracted)} characters from {path}")
    return extracted


def read_examples_as_text(path: str) -> str:
    """Convert example workbook into a calibration string for the system prompt."""
    df = pd.read_excel(path)

    required_like = [
        "Sample Question",
        "Point A", "Rationale",
        "Point B", "Rationale.1",
        "Point C", "Rationale.2",
        "Point D", "Rationale.3",
        "Point E", "Rationale.4",
        "Point F", "Rationale.5",
    ]
    missing = [c for c in required_like if c not in df.columns]
    if missing:
        raise ValueError(f"Examples file missing expected columns: {missing}")

    def safe_int(v):
        """Safely convert rubric example numeric fields."""
        if pd.isna(v):
            return 0
        if isinstance(v, str):
            v = v.strip()
            if v.lower() in ("1", "yes", "true", "x"):
                return 1
            if v.lower() in ("0", "no", "false", ""):
                return 0
            try:
                return int(float(v))
            except Exception:
                return 0
        try:
            return int(v)
        except Exception:
            return 0

    blocks = []
    for i, row in df.iterrows():
        try:
            a = safe_int(row["Point A"])
            b = safe_int(row["Point B"])
            c = safe_int(row["Point C"])
            d = safe_int(row["Point D"])
            e = safe_int(row["Point E"])
            f = safe_int(row["Point F"])
        except:
            raise ValueError(f"Examples row {i+1} has non-numeric values.")

        feedback_total = a + b + c + d
        systems_total = e + f

        block = f"""
EXAMPLE {i+1}
Response:
{str(row["Sample Question"]).strip()}

Human Scores:
A={a}, B={b}, C={c}, D={d}, E={e}, F={f}
Feedback_Total={feedback_total}, Systems_Total={systems_total}

Human Rationales:
A: {str(row["Rationale"]).strip()}
B: {str(row["Rationale.1"]).strip()}
C: {str(row["Rationale.2"]).strip()}
D: {str(row["Rationale.3"]).strip()}
E: {str(row["Rationale.4"]).strip()}
F: {str(row["Rationale.5"]).strip()}
""".strip()

        blocks.append(block)

    return "\n\n".join(blocks)

###################################################################################################
# Column Detection
###################################################################################################

def detect_response_column(df: pd.DataFrame) -> str:
    """Automatically determine the student response column."""
    preferred_patterns = [
        r"^Q2", r"response", r"answer", r"text", r"essay", r"written"
    ]

    for pattern in preferred_patterns:
        for col in df.columns:
            if re.search(pattern, col, flags=re.IGNORECASE):
                return col

    # Fall back to longest average string column — log a warning so this doesn't go unnoticed
    best_col = None
    best_len = -1
    for col in df.columns:
        avg_len = df[col].astype(str).str.len().mean()
        if avg_len > best_len:
            best_len = avg_len
            best_col = col

    logger.warning(
        f"No preferred response column found. Falling back to longest column: '{best_col}'. "
        f"Verify this is correct before trusting grades."
    )
    return best_col

###################################################################################################
# Prompt Construction
# FIX 4, 5, 6, 7, 8 all live here
###################################################################################################

def build_system_prompt(rubric_text: str, examples_text: str) -> str:
    """
    System instructions with calibration rules, inference guidance, and loop-type branching.
    """
    return f"""
You are an academic grading engine for Babson College.
You evaluate systemic reasoning using a strict rubric.

=============================================================
RUBRIC PREAMBLE
=============================================================
Reasonable inference may be applied when a logical outcome necessarily follows from the
stated relationships, even if not fully articulated by the student. Award points when a
student's meaning is clear from context, even if exact terminology is absent.

Point D requires that the consequence of the fish/fisher relationship feeds back to alter
future fish or fisher populations, completing the loop over time. Time language alone
("eventually", "over time") is not sufficient — the loop must actually close.

=============================================================
STEP-BY-STEP SCORING PROCEDURE
=============================================================

STEP 0 — CLASSIFY LOOP TYPE
Before scoring Points C and D, determine which type of feedback loop the student describes:

  BALANCING:    The relationship self-corrects toward stability.
                e.g., fewer fish -> fewer fishers -> fish population recovers.
  REINFORCING:  The relationship amplifies in one direction.
                e.g., fewer fish -> fishers work harder -> fish population continues to decline.
  UNCLEAR:      The response does not describe either loop clearly.

Record this as "Loop_Type": "balancing", "reinforcing", or "unclear".
Points C and D must be scored against the branch that matches Loop_Type.

---

STEP 1 — SCORE POINT A (same for both loop types)
Award A=1 if the response indicates ANY connection between fishers and fish.
This is a low bar — nearly any mention of fishers catching, depending on, or affecting
fish qualifies.

---

STEP 2 — SCORE POINT B (same for both loop types)
Award B=1 if the response describes how one population changes RELATIVE to the other.
This means a direct relationship (more fish -> more fishers) OR an inverse relationship
(more fishers -> fewer fish). Words like increase, decrease, more, less, attract, reduce,
or references to supply and demand all satisfy this requirement.

---

STEP 3 — SCORE POINT C (branch by Loop_Type)

  BALANCING branch:
    Award C=1 if the response describes a second relationship that is the opposite of B —
    i.e., the consequence of B feeding back in the other direction.
    Example: more fishing reduces fish -> smaller fish population leads to fewer fishers
    OR introduces some regulatory/corrective response.
    The loop is not yet established at C — that is Point D's job.

  REINFORCING branch:
    Award C=1 if the response describes a change in fisher behavior (increased effort,
    new technology, etc.) that maintains or increases catch despite a declining fish population.
    Again, the loop is not yet established at C.

---

STEP 4 — SCORE POINT D — LOOP CLOSURE TEST
Ask yourself two questions:
  1. Does the response describe an outcome that loops back to alter the original relationship?
  2. Does the loop actually close — does a consequence ultimately affect fish population,
     fisher population, or rate of fishing again?

  BALANCING branch:
    Award D=1 if the response establishes long-term STABILITY — either through a predator-prey
    cycle or through an intervention (e.g., regulation, fish farming) that maintains a viable
    population of both fish and fishers. The key signal is stability or recovery.

  REINFORCING branch:
    Award D=1 if the response establishes long-term CONSEQUENCES — typically collapse of the
    fishery, unsustainable overfishing, or eventual exit of fishers. The key signal is
    escalation to a conclusion.

  Do NOT award D for time language alone. The loop must close.
  D requires A=1, B=1, C=1.

---

STEP 5 — SCORE POINT E
Award E=1 if the response mentions system elements OUTSIDE the fish/fisher relationship,
such as: economic systems, government or regulation, broader ecosystems, or general human
consumption (food supply, markets, other species).

These phrases DO earn E:
  - "fishers rely on fish for their livelihood" (livelihood implies an economic system)
  - "fish are sold to consumers" or "sold for money" (broader market or food supply)
  - "overfishing can damage the ocean ecosystem" (broader ecological system)
  - "regulations may be introduced" (government or policy system)
  - "famine" or "food shortage" for a community (societal system)

These phrases do NOT earn E:
  - "fish act as a food source for fishers"
  - "fishers depend on fish for food for themselves"
  - "fishermen kill fish to eat them"
  These describe only the fishers' own consumption — they do not extend to a broader system.

---

STEP 6 — SCORE POINT F
Award F=1 ONLY if E=1, AND the response goes further to describe the IMPACT of the
fish/fisher relationship on the broader system, OR the impact of the broader system on
the fish/fisher relationship.
A mere mention earns E. Elaboration of impact earns F.
Examples of F-level elaboration:
  - "overfishing would raise the market price of fish for consumers"
  - "declining fish populations affect other species in the food chain"
  - "fish farming reduces demand for wild-caught fish, reducing the fisher labor force"

=============================================================
EVIDENCE STANDARD FOR JUSTIFICATIONS
=============================================================
Each justification must reference specific language or reasoning from the student's response.
You do NOT need a verbatim quote — paraphrase is acceptable — but your justification must
clearly identify the part of the response that earned or did not earn the point.

ACCEPTABLE (no quote required):
  "The student describes fishers catching fish and fish populations declining, which
   establishes the inverse relationship required for Point B."

NOT ACCEPTABLE (too vague):
  "The student addressed the relationship between fish and fishers."

=============================================================
SCORING CONSTRAINTS
=============================================================
- D requires A=1, B=1, C=1. Do not award D if any of A, B, or C is 0.
- F requires E=1. Do not award F if E=0.
- Totals must equal the sum of their sub-scores.
- Output must be valid JSON only — no markdown, no explanation outside the JSON.

=============================================================
RUBRIC
=============================================================
{rubric_text}

=============================================================
CALIBRATION EXAMPLES (human-scored)
=============================================================
{examples_text}
"""


def build_user_prompt(response_text: str, response_id: Any) -> str:
    """User grading prompt requesting exact JSON format including Loop_Type."""
    return f"""
Grade this student response:

Response_ID: {response_id}

Student Response:
\"\"\"
{response_text}
\"\"\"

Follow the step-by-step procedure from the system prompt.
Classify Loop_Type FIRST, then score A through F in order.

Return JSON ONLY with this exact structure:

{{
 "Loop_Type": "balancing" or "reinforcing" or "unclear",

 "Feedback_A_Score": 0 or 1,
 "Feedback_A_Justification": "reference specific student language or reasoning",

 "Feedback_B_Score": 0 or 1,
 "Feedback_B_Justification": "reference specific student language or reasoning",

 "Feedback_C_Score": 0 or 1,
 "Feedback_C_Justification": "reference specific student language or reasoning",

 "Feedback_D_Score": 0 or 1,
 "Feedback_D_Justification": "reference specific student language and explain how the loop closes",

 "Systems_E_Score": 0 or 1,
 "Systems_E_Justification": "reference specific student language or reasoning",

 "Systems_F_Score": 0 or 1,
 "Systems_F_Justification": "reference specific student language or reasoning",

 "Total_Feedback_Score": integer,
 "Total_Systems_Score": integer,
 "Overall_Total": integer,

 "Grading_Rationale": "one or two sentences summarizing scoring decisions and loop type classification"
}}
"""

###################################################################################################
# JSON Extraction and Repair
###################################################################################################

def extract_json(text: str) -> Dict[str, Any]:
    """Safely extract JSON dictionary from model output."""
    text = text.strip()

    # Remove markdown fencing
    if text.startswith("```"):
        text = re.sub(r"^```(?:json)?", "", text).strip()
        text = re.sub(r"```$", "", text).strip()

    # Try direct JSON load
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        pass

    # Fallback: extract first {...} block
    match = re.search(r"\{.*?\}", text, flags=re.DOTALL)
    if match:
        try:
            return json.loads(match.group(0))
        except:
            pass

    raise ValueError("Could not decode JSON from model output.")

###################################################################################################
# VALIDATION + AUTO-REPAIR ENGINE
###################################################################################################

def validate_and_repair(result: Dict[str, Any], response_text: str) -> Dict[str, Any]:
    """Smart repair layer preventing hard failures."""
    REQUIRED = [
        "Loop_Type",
        "Feedback_A_Score", "Feedback_A_Justification",
        "Feedback_B_Score", "Feedback_B_Justification",
        "Feedback_C_Score", "Feedback_C_Justification",
        "Feedback_D_Score", "Feedback_D_Justification",
        "Systems_E_Score", "Systems_E_Justification",
        "Systems_F_Score", "Systems_F_Justification",
        "Total_Feedback_Score", "Total_Systems_Score",
        "Overall_Total", "Grading_Rationale"
    ]

    # 1. Ensure all required keys exist
    for key in REQUIRED:
        if key not in result:
            if key == "Loop_Type":
                result[key] = "unclear"
            elif key.endswith("_Score"):
                result[key] = 0
            elif key.endswith("_Justification"):
                result[key] = "No justification provided."
            else:
                result[key] = ""

    # 2. Normalize Loop_Type to valid values
    if result["Loop_Type"] not in ("balancing", "reinforcing", "unclear"):
        result["Loop_Type"] = "unclear"

    # 3. Normalize scores to binary
    for k in [x for x in result if x.endswith("_Score")]:
        if result[k] not in (0, 1):
            result[k] = 0

    # 4. Enforce D requires A=B=C=1
    if result["Feedback_D_Score"] == 1:
        if not (result["Feedback_A_Score"] and result["Feedback_B_Score"] and result["Feedback_C_Score"]):
            result["Feedback_D_Score"] = 0

    # 5. Enforce F requires E
    if result["Systems_F_Score"] == 1 and result["Systems_E_Score"] == 0:
        result["Systems_F_Score"] = 0

    # FIX 3: Fabricated quote injection removed.
    # Original code silently prepended first 12 words of the student response to any
    # justification lacking a quote mark, creating misleading audit trails.
    # Missing or empty justifications are now flagged for human review instead.
    for key in [k for k in result if k.endswith("_Justification")]:
        text = str(result[key]).strip()
        if not text or text == "No justification provided.":
            logger.warning(f"Empty justification for {key} — flagging for review.")
            result[key] = "[REVIEW NEEDED: justification missing]"

    # 6. Recompute totals
    fb_total = (
        result["Feedback_A_Score"]
        + result["Feedback_B_Score"]
        + result["Feedback_C_Score"]
        + result["Feedback_D_Score"]
    )
    sys_total = result["Systems_E_Score"] + result["Systems_F_Score"]
    overall   = fb_total + sys_total

    result["Total_Feedback_Score"] = fb_total
    result["Total_Systems_Score"]  = sys_total
    result["Overall_Total"]        = overall

    return result

###################################################################################################
# FIX 9: BORDERLINE SCORE FLAGGING
###################################################################################################

BORDERLINE_FEEDBACK_TOTALS = {2, 3}  # 2<->3 and 3<->4 boundaries
BORDERLINE_SYSTEMS_TOTALS  = {1}     # E awarded but not F — worth checking F eligibility

def flag_borderline(result: Dict[str, Any], response_id: Any) -> str:
    """
    Return a human-readable flag string if this result falls in a borderline range.
    Returns empty string if no review is needed.

    Borderline conditions:
      - Feedback total of 2 or 3 (boundary between C/D award decisions)
      - Systems total of 1 (E awarded but not F — worth checking F eligibility)
      - Loop_Type of 'unclear' (model couldn't classify — highest ambiguity)
    """
    reasons = []

    if result["Total_Feedback_Score"] in BORDERLINE_FEEDBACK_TOTALS:
        reasons.append(f"Feedback={result['Total_Feedback_Score']} (boundary)")

    if result["Total_Systems_Score"] in BORDERLINE_SYSTEMS_TOTALS:
        reasons.append("Systems=1 (check F eligibility)")

    if result.get("Loop_Type") == "unclear":
        reasons.append("Loop_Type=unclear (ambiguous response)")

    if reasons:
        flag = "; ".join(reasons)
        logger.warning(f"[BORDERLINE] Response_ID {response_id}: {flag}")
        return flag

    return ""

###################################################################################################
# GRADING ENGINE (with retry + repair)
###################################################################################################

@retry(stop=stop_after_attempt(4), wait=wait_exponential(multiplier=1, min=2, max=20))
def grade_one(client, model, system_prompt, response_text, response_id):
    """Main grading call with built-in repair + retry."""
    user_prompt = build_user_prompt(response_text, response_id)

    try:
        resp = client.chat.completions.create(
            model=model,
            messages=[
                {"role": "system", "content": system_prompt},
                {"role": "user", "content": user_prompt}
            ],
            max_tokens=1800,
            temperature=0,  # FIX 2: was 0.1 — deterministic grading requires 0
        )
        raw = resp.choices[0].message.content

        # First extraction attempt
        try:
            data = extract_json(raw)
        except Exception:
            # Repair attempt
            repair_prompt = f"""
The following model output is INVALID JSON.
Fix it and return VALID JSON ONLY.

Original output:
{raw}
"""
            repair = client.chat.completions.create(
                model=model,
                messages=[
                    {"role": "system", "content": "You fix JSON formatting errors."},
                    {"role": "user", "content": repair_prompt}
                ],
                max_tokens=1800,
                temperature=0,
            )
            data = extract_json(repair.choices[0].message.content)

        # Apply repair / validation rules
        final = validate_and_repair(data, response_text)
        return final

    except Exception as e:
        logger.error(f"[ERROR] Response_ID {response_id}: {e}")
        raise

###################################################################################################
# Utility: Save checkpoint
###################################################################################################
import time
import os

def safe_save_excel(df, path, max_retries=5):
    """
    Saves the DataFrame to Excel with automatic detection of file locks.
    If the file is locked (PermissionError), it retries or saves to an alternate filename.
    """
    attempt = 1
    base, ext = os.path.splitext(path)

    while attempt <= max_retries:
        try:
            df.to_excel(path, index=False, engine="openpyxl")
            print(f"💾 Saved: {path}")
            return path

        except PermissionError:
            print(f"⚠️ File is locked (attempt {attempt}/{max_retries}): {path}")
            time.sleep(1)
            attempt += 1

    fallback_path = base + "_UNLOCKED_COPY" + ext
    print(f"🚨 Unable to save to locked file.\n➡️ Saving instead to: {fallback_path}")

    df.to_excel(fallback_path, index=False, engine="openpyxl")
    return fallback_path

def save_checkpoint(df: pd.DataFrame, path: str):
    saved_path = safe_save_excel(df, path)
    logger.info(f"Checkpoint saved → {saved_path}")

###################################################################################################
# MAIN EXECUTION LOOP
###################################################################################################

def main():
    print("\n🤖 Babson College AI Grading Assistant v3")
    print("=" * 60)
    print(f"Started: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n")

    default_input    = "Master_without_Identifiers_5.22.25.xlsx"
    default_rubric   = "Revised_Rubric_Aug2025.docx"
    default_examples = "SampleRubric_AI_SES.xlsx"
    default_output   = "AI_Innovation_Fellowship_Graded_Output.xlsx"

    class Args:
        def __init__(self):
            self.input    = default_input
            self.rubric   = default_rubric
            self.examples = default_examples
            self.output   = default_output
            self.model    = DEFAULT_MODEL
            self.limit    = None

    args = Args()

    # Resolve real file paths
    script_dir    = os.path.dirname(os.path.abspath(__file__))
    args.input    = os.path.join(script_dir, args.input)
    args.rubric   = os.path.join(script_dir, args.rubric)
    args.examples = os.path.join(script_dir, args.examples)
    args.output   = os.path.join(script_dir, args.output)

    # Environment
    endpoint    = os.getenv("AZURE_OPENAI_ENDPOINT")
    api_key     = os.getenv("AZURE_OPENAI_API_KEY")
    deployment  = os.getenv("AZURE_OPENAI_DEPLOYMENT")
    api_version = os.getenv("AZURE_OPENAI_API_VERSION")

    if not all([endpoint, api_key, deployment, api_version]):
        raise EnvironmentError("Missing required Azure environment variables.")

    # Azure Client
    client = AzureOpenAI(
        azure_endpoint=endpoint,
        api_key=api_key,
        api_version=api_version,
    )

    # Load rubric + examples
    rubric_text   = read_docx_text(args.rubric)
    examples_text = read_examples_as_text(args.examples)
    system_prompt = build_system_prompt(rubric_text, examples_text)

    # Diagnostic: verify rubric table content was captured
    print(f"📄 Rubric loaded: {len(rubric_text):,} characters")
    if len(rubric_text) < 500:
        print("⚠️  WARNING: Rubric text is very short — table content may not have been captured.")
        print("   Verify the .docx file path and structure before proceeding.")

    # Load spreadsheet
    df = pd.read_excel(args.input)
    response_col = detect_response_column(df)

    if "Response_ID" not in df.columns:
        df["Response_ID"] = range(1, len(df) + 1)

    grading_columns = [
        "Loop_Type",
        "Feedback_A_Score", "Feedback_A_Justification",
        "Feedback_B_Score", "Feedback_B_Justification",
        "Feedback_C_Score", "Feedback_C_Justification",
        "Feedback_D_Score", "Feedback_D_Justification",
        "Systems_E_Score", "Systems_E_Justification",
        "Systems_F_Score", "Systems_F_Justification",
        "Total_Feedback_Score", "Total_Systems_Score",
        "Overall_Total", "Grading_Rationale",
        "Needs_Review",  # FIX 9: borderline flag
    ]
    for col in grading_columns:
        if col not in df.columns:
            df[col] = None

    # Grading loop
    total = len(df)
    print(f"📊 Total responses: {total}")
    print(f"📝 Response column detected: {response_col}")
    print("🚀 Grading begins:")
    print("-" * 60)

    start_time = time.time()
    successful = 0
    failed     = 0
    borderline = 0

    for i in range(total):
        rid  = df.at[i, "Response_ID"]
        text = str(df.at[i, response_col]).strip()

        if not text or text.lower() == "nan":
            logger.warning(f"Skipping empty row {i} (Response_ID={rid})")
            continue

        try:
            print(f"🔄 Grading row {i+1}/{total} (Response_ID={rid})")
            result = grade_one(
                client=client,
                model=deployment,
                system_prompt=system_prompt,
                response_text=text,
                response_id=rid,
            )

            # FIX 9: Check for borderline scores and record flag
            review_flag = flag_borderline(result, rid)
            result["Needs_Review"] = review_flag if review_flag else ""
            if review_flag:
                borderline += 1
                print(f"🔍 Row {i+1} flagged for review: {review_flag}")

            for k, v in result.items():
                df.at[i, k] = v

            successful += 1
            print(f"✅ Completed row {i+1}/{total} — Score: {result['Overall_Total']}/6"
                  f" [{result.get('Loop_Type', '?')}]")

        except Exception as e:
            failed += 1
            msg = f"Failed to grade row {i} (Response_ID={rid}): {e}"
            df.at[i, "Grading_Rationale"] = f"ERROR: {e}"
            logger.error(msg)
            print(f"❌ {msg}")

        # Save checkpoints every 5 rows
        if (i + 1) % 5 == 0:
            save_checkpoint(df, args.output.replace(".xlsx", "_checkpoint.xlsx"))
            print("💾 Checkpoint saved.")

        time.sleep(0.3)

    # Final save
    safe_save_excel(df, args.output)

    total_time = time.time() - start_time
    print("\n" + "=" * 60)
    print("🎉 GRADING COMPLETE")
    print("=" * 60)
    print(f"✅ Success:     {successful}")
    print(f"❌ Failed:      {failed}")
    print(f"🔍 Borderline:  {borderline} (see Needs_Review column)")
    print(f"⏱️  Time:        {total_time:.1f} seconds")
    print(f"💾 Saved:       {args.output}")

###################################################################################################
# ENTRY POINT
###################################################################################################

if __name__ == "__main__":
    main()
