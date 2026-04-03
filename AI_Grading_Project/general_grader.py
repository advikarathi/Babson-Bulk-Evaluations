###################################################################################################
# AI RUBRIC GRADER FOR VS CODE
#
# What this version supports
# - Rubric files: PDF, DOCX, TXT, MD, XLSX, XLS, CSV
# - Response files: XLSX, XLS, CSV
# - Uses PyMuPDF for PDF extraction, which is usually more reliable than pdfplumber
# - Works well in VS Code with .env, relative paths, argparse, and interactive fallback prompts
# - Generates a rubric schema first, then grades responses against that schema
# - Exports:
#     1. graded output Excel
#     2. rubric schema JSON
#     3. run summary JSON
#     4. checkpoint Excel while grading
# - Includes dependency enforcement, JSON repair, empty response handling, and review flags
#
# Honest note
# No grader is "the best in the world" for every rubric. This is a very strong generalized grader,
# but you should still review schema extraction and borderline cases before using it at scale.
###################################################################################################

import argparse
import json
import logging
import os
import re
import sys
import time
from dataclasses import dataclass, field, asdict
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

try:
    import fitz  # PyMuPDF
except ImportError:
    raise ImportError(
        "PyMuPDF is not installed. Run: pip install pymupdf"
    )
import pandas as pd
import docx
from dotenv import load_dotenv
from openai import AzureOpenAI
from tenacity import retry, stop_after_attempt, wait_exponential


###################################################################################################
# VS CODE / WORKSPACE SETUP
###################################################################################################

SCRIPT_DIR = Path(__file__).resolve().parent
load_dotenv(SCRIPT_DIR / ".env")

LOG_FILE = SCRIPT_DIR / "grading_log.txt"

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s | %(levelname)s | %(message)s",
    handlers=[
        logging.FileHandler(LOG_FILE, mode="a", encoding="utf-8"),
        logging.StreamHandler()
    ]
)
logger = logging.getLogger(__name__)

DEFAULT_MODEL = os.getenv("AZURE_OPENAI_DEPLOYMENT", "gpt-4.1-mini")
DEFAULT_API_VERSION = os.getenv("AZURE_OPENAI_API_VERSION", "2024-02-15-preview")


###################################################################################################
# DATA CLASSES
###################################################################################################

@dataclass
class RubricCriterion:
    id: str
    name: str
    description: str
    max_score: int = 1
    dependencies: List[str] = field(default_factory=list)
    scoring_notes: List[str] = field(default_factory=list)
    is_binary: bool = True


@dataclass
class ParsedRubric:
    title: str
    criteria: List[RubricCriterion]
    raw_text: str
    scoring_notes: str = ""
    source_path: str = ""

    def get_criterion_ids(self) -> List[str]:
        return [c.id for c in self.criteria]

    def get_max_total(self) -> int:
        return sum(c.max_score for c in self.criteria)


@dataclass
class GradingConfig:
    rubric_path: str
    responses_path: str
    output_path: str
    response_column: Optional[str] = None
    id_column: Optional[str] = None
    rubric_sheet_name: Optional[str] = None
    responses_sheet_name: Optional[str] = None
    model: str = DEFAULT_MODEL
    checkpoint_interval: int = 5
    limit: Optional[int] = None
    preview_only: bool = False
    export_schema_path: Optional[str] = None
    export_summary_path: Optional[str] = None


###################################################################################################
# FILE READING UTILITIES
###################################################################################################

def resolve_path(path_str: str) -> Path:
    path = Path(path_str).expanduser()
    if not path.is_absolute():
        path = (SCRIPT_DIR / path).resolve()
    return path


# DOCX

def read_docx_text(path: Path) -> str:
    doc = docx.Document(str(path))
    parts: List[str] = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            parts.append(text)

    for table in doc.tables:
        for row in table.rows:
            row_parts = []
            for cell in row.cells:
                cell_text = cell.text.strip()
                if cell_text:
                    row_parts.append(cell_text)
            if row_parts:
                parts.append(" | ".join(row_parts))
        parts.append("")

    extracted = "\n".join(parts).strip()
    if not extracted:
        raise ValueError(f"DOCX appears empty: {path}")
    return extracted


# TXT / MD

def read_text_file(path: Path) -> str:
    with open(path, "r", encoding="utf-8") as f:
        text = f.read().strip()
    if not text:
        raise ValueError(f"Text file appears empty: {path}")
    return text


# PDF via PyMuPDF

def read_pdf_text(path: Path) -> str:
    parts: List[str] = []
    doc = fitz.open(str(path))

    try:
        for i, page in enumerate(doc, start=1):
            try:
                text = page.get_text("text")
                if text and text.strip():
                    parts.append(text.strip())
            except Exception as e:
                logger.warning(f"Could not extract text from PDF page {i}: {e}")
    finally:
        doc.close()

    extracted = "\n\n".join(parts).strip()
    if not extracted:
        raise ValueError(
            f"No extractable text found in PDF: {path}. It may be scanned/image based."
        )
    return extracted


# Spreadsheet to text for rubric support

def read_spreadsheet_as_text(path: Path, sheet_name: Optional[str] = None) -> str:
    if path.suffix.lower() == ".csv":
        df = pd.read_csv(path)
    else:
        df = pd.read_excel(path, sheet_name=sheet_name or 0)

    if df.empty:
        raise ValueError(f"Spreadsheet appears empty: {path}")

    parts: List[str] = []
    parts.append("COLUMNS: " + ", ".join([str(c) for c in df.columns]))

    for _, row in df.iterrows():
        row_parts = []
        for col in df.columns:
            value = row[col]
            if pd.notna(value) and str(value).strip() != "":
                row_parts.append(f"{col}: {str(value).strip()}")
        if row_parts:
            parts.append(" | ".join(row_parts))

    extracted = "\n".join(parts).strip()
    if not extracted:
        raise ValueError(f"Could not extract rubric text from spreadsheet: {path}")
    return extracted


# Main rubric reader

def read_rubric_file(path: Path, sheet_name: Optional[str] = None) -> str:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return read_docx_text(path)
    if suffix in [".txt", ".md"]:
        return read_text_file(path)
    if suffix == ".pdf":
        return read_pdf_text(path)
    if suffix in [".xlsx", ".xls", ".csv"]:
        return read_spreadsheet_as_text(path, sheet_name=sheet_name)

    raise ValueError(
        f"Unsupported rubric format: {path.name}. Supported: .pdf, .docx, .txt, .md, .xlsx, .xls, .csv"
    )


# Response reader

def read_responses_file(path: Path, sheet_name: Optional[str] = None) -> pd.DataFrame:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return pd.read_csv(path)
    if suffix in [".xlsx", ".xls"]:
        return pd.read_excel(path, sheet_name=sheet_name or 0)
    raise ValueError(f"Unsupported responses format: {path.name}. Supported: .xlsx, .xls, .csv")


###################################################################################################
# OPENAI / AZURE CLIENT
###################################################################################################

def get_client() -> AzureOpenAI:
    endpoint = os.getenv("AZURE_OPENAI_ENDPOINT")
    api_key = os.getenv("AZURE_OPENAI_API_KEY")
    api_version = os.getenv("AZURE_OPENAI_API_VERSION", DEFAULT_API_VERSION)

    if not endpoint or not api_key:
        raise EnvironmentError(
            "Missing Azure OpenAI credentials. Please set AZURE_OPENAI_ENDPOINT and AZURE_OPENAI_API_KEY in your .env file."
        )

    return AzureOpenAI(
        azure_endpoint=endpoint,
        api_key=api_key,
        api_version=api_version,
    )


###################################################################################################
# COLUMN DETECTION
###################################################################################################

def detect_response_column(df: pd.DataFrame, hint: Optional[str] = None) -> str:
    if hint and hint in df.columns:
        return hint

    patterns = [
        r"response", r"answer", r"essay", r"submission", r"text", r"written",
        r"comment", r"discussion", r"content", r"body"
    ]

    for pattern in patterns:
        for col in df.columns:
            if re.search(pattern, str(col), flags=re.IGNORECASE):
                return str(col)

    best_col = None
    best_len = -1
    for col in df.columns:
        avg_len = df[col].astype(str).str.len().mean()
        if avg_len > best_len:
            best_len = avg_len
            best_col = str(col)

    if best_col is None:
        raise ValueError("Could not detect response column.")

    logger.warning(f"No obvious response column found. Using longest average text column: {best_col}")
    return best_col



def detect_id_column(df: pd.DataFrame, hint: Optional[str] = None) -> Optional[str]:
    if hint and hint in df.columns:
        return hint

    patterns = [r"response_id", r"student_id", r"submission_id", r"id", r"number", r"name"]
    for pattern in patterns:
        for col in df.columns:
            if re.search(pattern, str(col), flags=re.IGNORECASE):
                return str(col)
    return None


###################################################################################################
# JSON UTILITIES
###################################################################################################

def extract_json(text: str) -> Dict[str, Any]:
    text = text.strip()

    if text.startswith("```"):
        text = re.sub(r"^```(?:json)?", "", text).strip()
        text = re.sub(r"```$", "", text).strip()

    try:
        return json.loads(text)
    except json.JSONDecodeError:
        pass

    match = re.search(r"\{.*\}", text, flags=re.DOTALL)
    if match:
        try:
            return json.loads(match.group(0))
        except Exception:
            pass

    raise ValueError("Could not decode JSON from model output.")


###################################################################################################
# RUBRIC PARSING WITH AI
###################################################################################################

def validate_parsed_rubric(rubric: ParsedRubric) -> ParsedRubric:
    if not rubric.criteria:
        raise ValueError("Rubric parsing produced zero criteria.")

    seen = set()
    for c in rubric.criteria:
        c.id = str(c.id).strip()
        c.name = str(c.name).strip() or f"Criterion {c.id}"
        c.description = str(c.description).strip() or c.name

        if not c.id:
            raise ValueError("Found rubric criterion with empty id.")
        if c.id in seen:
            raise ValueError(f"Duplicate criterion id found: {c.id}")
        seen.add(c.id)

        try:
            c.max_score = int(c.max_score)
        except Exception:
            c.max_score = 1
        if c.max_score < 0:
            c.max_score = 0

        if not isinstance(c.dependencies, list):
            c.dependencies = []
        c.dependencies = [str(x).strip() for x in c.dependencies if str(x).strip()]

        if not isinstance(c.scoring_notes, list):
            c.scoring_notes = []

        for dep in c.dependencies:
            if dep == c.id:
                raise ValueError(f"Criterion {c.id} cannot depend on itself.")

    criterion_ids = set(rubric.get_criterion_ids())
    for c in rubric.criteria:
        c.dependencies = [dep for dep in c.dependencies if dep in criterion_ids]

    return rubric



def build_rubric_parse_prompt(rubric_text: str) -> str:
    template = (
        'Analyze this grading rubric and convert it into a machine readable grading structure.\n\n'
        'Return valid JSON only with this exact structure:\n'
        '{\n'
        '  "title": "rubric title",\n'
        '  "scoring_notes": "overall notes for applying the rubric",\n'
        '  "criteria": [\n'
        '    {\n'
        '      "id": "original label like A or 1 or thesis",\n'
        '      "name": "short human readable criterion name",\n'
        '      "description": "full description of what earns points",\n'
        '      "max_score": 1,\n'
        '      "dependencies": ["ids that must already be met"],\n'
        '      "scoring_notes": ["specific notes for this criterion"],\n'
        '      "is_binary": true\n'
        '    }\n'
        '  ]\n'
        '}\n\n'
        'Rules:\n'
        '1. Extract each distinct scoring criterion separately.\n'
        '2. Preserve the rubric\'s real structure. Do not invent criteria.\n'
        '3. If the rubric has scoring bands, convert each criterion into a max_score that matches the rubric.\n'
        '4. If max_score is not specified, use 1.\n'
        '5. If a criterion clearly requires another, include that in dependencies.\n'
        '6. Keep the rubric descriptions detailed enough for grading.\n'
        '7. If the rubric includes overall instructions, place them in scoring_notes.\n'
        '8. Return JSON only.\n\n'
        'RUBRIC TEXT:\n'
        '"""\n'
        + rubric_text +
        '\n"""'
    )
    return template

@retry(stop=stop_after_attempt(4), wait=wait_exponential(multiplier=1, min=2, max=20))
def parse_rubric_with_ai(client: AzureOpenAI, model: str, rubric_text: str, source_path: str = "") -> ParsedRubric:
    prompt = build_rubric_parse_prompt(rubric_text)

    response = client.chat.completions.create(
        model=model,
        messages=[
            {"role": "system", "content": "You are a rubric parsing assistant. Extract grading structures accurately and conservatively."},
            {"role": "user", "content": prompt}
        ],
        max_tokens=4000,
        temperature=0,
    )

    raw = response.choices[0].message.content

    try:
        data = extract_json(raw)
    except Exception:
        repair = client.chat.completions.create(
            model=model,
            messages=[
                {"role": "system", "content": "You repair broken JSON only."},
                {"role": "user", "content": f"Fix this into valid JSON only:\n\n{raw}"}
            ],
            max_tokens=4000,
            temperature=0,
        )
        data = extract_json(repair.choices[0].message.content)

    criteria: List[RubricCriterion] = []
    for c in data.get("criteria", []):
        criteria.append(RubricCriterion(
            id=str(c.get("id", "")).strip(),
            name=str(c.get("name", "")).strip(),
            description=str(c.get("description", "")).strip(),
            max_score=c.get("max_score", 1),
            dependencies=c.get("dependencies", []),
            scoring_notes=c.get("scoring_notes", []),
            is_binary=bool(c.get("is_binary", True))
        ))

    rubric = ParsedRubric(
        title=str(data.get("title", "Untitled Rubric")).strip(),
        criteria=criteria,
        raw_text=rubric_text,
        scoring_notes=str(data.get("scoring_notes", "")).strip(),
        source_path=source_path,
    )

    return validate_parsed_rubric(rubric)


###################################################################################################
# SCHEMA EXPORT
###################################################################################################

def rubric_to_dict(rubric: ParsedRubric) -> Dict[str, Any]:
    return {
        "title": rubric.title,
        "source_path": rubric.source_path,
        "scoring_notes": rubric.scoring_notes,
        "total_max_score": rubric.get_max_total(),
        "criteria": [asdict(c) for c in rubric.criteria],
    }



def save_json(data: Dict[str, Any], path: Path) -> None:
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2, ensure_ascii=False)


###################################################################################################
# DYNAMIC PROMPT GENERATION
###################################################################################################

def build_dynamic_system_prompt(rubric: ParsedRubric) -> str:
    criteria_text = []
    for c in rubric.criteria:
        deps_text = f"\nDependencies: {', '.join(c.dependencies)}" if c.dependencies else ""
        extra_notes = f"\nNotes: {'; '.join(c.scoring_notes)}" if c.scoring_notes else ""
        criteria_text.append(
            f"CRITERION {c.id}: {c.name}\n"
            f"Max Score: {c.max_score}\n"
            f"Description: {c.description}{deps_text}{extra_notes}"
        )

    output_schema = {
        "response_id": "original response id",
        "criteria_results": {
            c.id: {
                "score": f"integer from 0 to {c.max_score}",
                "justification": "specific evidence based explanation"
            } for c in rubric.criteria
        },
        "total_score": "integer",
        "max_possible_score": rubric.get_max_total(),
        "overall_rationale": "1 to 3 sentence summary"
    }

    return f"""
You are a strict academic grading engine.
Grade student responses only using the rubric below.

RUBRIC TITLE:
{rubric.title}

RUBRIC LEVEL NOTES:
{rubric.scoring_notes}

CRITERIA:
{chr(10).join(criteria_text)}

SCORING RULES:
1. Use only the rubric. Do not invent standards.
2. Respect dependencies. If a dependency is not met, dependent criteria must not receive points.
3. Every criterion must include a score and justification.
4. Justifications must reference specific student language or reasoning.
5. Keep scores within each criterion's allowed range.
6. Return valid JSON only.
7. Recompute total_score accurately.
8. If the response is weak, missing, or off topic, score accordingly rather than guessing generously.

OUTPUT SHAPE:
{json.dumps(output_schema, indent=2, ensure_ascii=False)}
"""



def build_dynamic_user_prompt(response_text: str, response_id: Any) -> str:
    return f"""Grade this student response.

Response_ID: {response_id}

Student Response:
\"\"\"
{response_text}
\"\"\"

Return JSON only.
"""


###################################################################################################
# VALIDATION AND REPAIR OF GRADING OUTPUT
###################################################################################################

def validate_and_repair_dynamic(result: Dict[str, Any], rubric: ParsedRubric) -> Dict[str, Any]:
    if "criteria_results" not in result or not isinstance(result["criteria_results"], dict):
        result["criteria_results"] = {}

    criteria_map = {c.id: c for c in rubric.criteria}

    for cid, criterion in criteria_map.items():
        if cid not in result["criteria_results"] or not isinstance(result["criteria_results"][cid], dict):
            result["criteria_results"][cid] = {}

        score = result["criteria_results"][cid].get("score", 0)
        justification = str(result["criteria_results"][cid].get("justification", "")).strip()

        try:
            score = int(score)
        except Exception:
            score = 0

        score = max(0, min(score, criterion.max_score))

        if criterion.dependencies and score > 0:
            unmet = False
            for dep in criterion.dependencies:
                dep_score = result["criteria_results"].get(dep, {}).get("score", 0)
                try:
                    dep_score = int(dep_score)
                except Exception:
                    dep_score = 0
                if dep_score <= 0:
                    unmet = True
                    break
            if unmet:
                score = 0
                if justification:
                    justification += " Dependency rule applied."
                else:
                    justification = "[REVIEW NEEDED: dependency rule applied]"

        if not justification:
            justification = "[REVIEW NEEDED: justification missing]"

        result["criteria_results"][cid] = {
            "score": score,
            "justification": justification
        }

    total = sum(result["criteria_results"][cid]["score"] for cid in criteria_map)
    result["total_score"] = total
    result["max_possible_score"] = rubric.get_max_total()

    if not str(result.get("overall_rationale", "")).strip():
        result["overall_rationale"] = "[REVIEW NEEDED: rationale missing]"

    return result



def detect_borderline(result: Dict[str, Any], rubric: ParsedRubric) -> str:
    reasons: List[str] = []
    total = int(result.get("total_score", 0))
    max_total = rubric.get_max_total()

    if max_total > 2:
        midpoint = max_total / 2
        if abs(total - midpoint) <= 1:
            reasons.append(f"Total score {total}/{max_total} is near midpoint")

    if total == 0:
        reasons.append("Zero score")
    if total == max_total:
        reasons.append("Perfect score")

    for c in rubric.criteria:
        just = result["criteria_results"][c.id]["justification"]
        if "REVIEW NEEDED" in just:
            reasons.append(f"Criterion {c.id} missing strong justification")

    return "; ".join(reasons)


###################################################################################################
# GRADING ENGINE
###################################################################################################

@retry(stop=stop_after_attempt(4), wait=wait_exponential(multiplier=1, min=2, max=20))
def grade_one_response(
    client: AzureOpenAI,
    model: str,
    system_prompt: str,
    rubric: ParsedRubric,
    response_text: str,
    response_id: Any,
) -> Dict[str, Any]:
    user_prompt = build_dynamic_user_prompt(response_text, response_id)

    response = client.chat.completions.create(
        model=model,
        messages=[
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt}
        ],
        max_tokens=2500,
        temperature=0,
    )

    raw = response.choices[0].message.content

    try:
        data = extract_json(raw)
    except Exception:
        repair_prompt = f"Fix this into valid JSON only:\n\n{raw}"
        repair = client.chat.completions.create(
            model=model,
            messages=[
                {"role": "system", "content": "You repair broken JSON only."},
                {"role": "user", "content": repair_prompt}
            ],
            max_tokens=2500,
            temperature=0,
        )
        data = extract_json(repair.choices[0].message.content)

    return validate_and_repair_dynamic(data, rubric)


###################################################################################################
# OUTPUT SHAPING
###################################################################################################

def sanitize_column_name(name: str) -> str:
    cleaned = re.sub(r"[^A-Za-z0-9]+", "_", name.strip())
    cleaned = re.sub(r"_+", "_", cleaned).strip("_")
    return cleaned[:80] if cleaned else "Criterion"



def prepare_output_columns(df: pd.DataFrame, rubric: ParsedRubric) -> Tuple[pd.DataFrame, Dict[str, Tuple[str, str]]]:
    col_map: Dict[str, Tuple[str, str]] = {}

    for c in rubric.criteria:
        base = sanitize_column_name(f"{c.id}_{c.name}")
        score_col = f"{base}_Score"
        just_col = f"{base}_Justification"
        col_map[c.id] = (score_col, just_col)

        if score_col not in df.columns:
            df[score_col] = None
        if just_col not in df.columns:
            df[just_col] = None

    for col in ["Total_Score", "Max_Possible_Score", "Overall_Rationale", "Needs_Review"]:
        if col not in df.columns:
            df[col] = None

    return df, col_map


###################################################################################################
# SAVE UTILITIES
###################################################################################################

def safe_save_excel(df: pd.DataFrame, path: Path, max_retries: int = 5) -> Path:
    attempt = 1
    while attempt <= max_retries:
        try:
            df.to_excel(path, index=False, engine="openpyxl")
            logger.info(f"Saved file: {path}")
            return path
        except PermissionError:
            logger.warning(f"File locked on attempt {attempt}: {path}")
            time.sleep(1)
            attempt += 1

    fallback = path.with_name(path.stem + "_UNLOCKED_COPY" + path.suffix)
    df.to_excel(fallback, index=False, engine="openpyxl")
    logger.warning(f"Saved fallback file instead: {fallback}")
    return fallback


###################################################################################################
# INTERACTIVE CLI FOR VS CODE TERMINAL
###################################################################################################

def get_user_input(prompt: str, default: Optional[str] = None, required: bool = True) -> str:
    display = f"{prompt} [{default}]: " if default else f"{prompt}: "
    value = input(display).strip()

    if not value and default:
        return default
    if not value and required:
        print("This field is required.")
        return get_user_input(prompt, default, required)
    return value



def validate_file_exists(path_str: str, label: str) -> bool:
    path = resolve_path(path_str)
    if not path.exists():
        print(f"❌ {label} not found: {path}")
        return False
    return True



def list_columns(df: pd.DataFrame) -> None:
    print("\nAvailable columns:")
    for i, col in enumerate(df.columns, start=1):
        sample = ""
        try:
            if len(df) > 0:
                sample = str(df[col].iloc[0])[:60]
        except Exception:
            pass
        print(f"  {i}. {col} (sample: {sample}...)")



def interactive_setup() -> GradingConfig:
    print("\n" + "=" * 70)
    print("AI RUBRIC GRADER FOR VS CODE - SETUP")
    print("=" * 70)

    while True:
        rubric_path = get_user_input("Path to rubric file (.pdf, .docx, .txt, .md, .xlsx, .xls, .csv)")
        if validate_file_exists(rubric_path, "Rubric"):
            break

    while True:
        responses_path = get_user_input("Path to responses file (.xlsx, .xls, .csv)")
        if validate_file_exists(responses_path, "Responses"):
            break

    responses_df = read_responses_file(resolve_path(responses_path))
    list_columns(responses_df)

    detected_response = None
    try:
        detected_response = detect_response_column(responses_df)
    except Exception:
        pass

    response_column = get_user_input("Response column name", default=detected_response, required=True)
    if response_column not in responses_df.columns:
        try:
            idx = int(response_column) - 1
            response_column = str(responses_df.columns[idx])
        except Exception:
            raise ValueError(f"Response column not found: {response_column}")

    detected_id = detect_id_column(responses_df)
    id_column = get_user_input("ID column name, or press Enter to auto generate", default=detected_id, required=False)
    if id_column and id_column not in responses_df.columns:
        try:
            idx = int(id_column) - 1
            id_column = str(responses_df.columns[idx])
        except Exception:
            id_column = None

    default_output = str(resolve_path(responses_path).with_name(resolve_path(responses_path).stem + "_graded.xlsx"))
    output_path = get_user_input("Output file path", default=default_output)

    default_schema = str(Path(output_path).with_name(Path(output_path).stem + "_rubric_schema.json"))
    export_schema_path = get_user_input("Rubric schema JSON path", default=default_schema, required=False)

    default_summary = str(Path(output_path).with_name(Path(output_path).stem + "_run_summary.json"))
    export_summary_path = get_user_input("Run summary JSON path", default=default_summary, required=False)

    env_model = os.getenv("AZURE_OPENAI_DEPLOYMENT", DEFAULT_MODEL)
    model = get_user_input("Azure deployment name", default=env_model)

    limit_str = get_user_input("Optional test row limit, or press Enter for all rows", default="", required=False)
    limit = int(limit_str) if limit_str.strip() else None

    preview_answer = get_user_input("Preview rubric schema only without grading? y/n", default="n")
    preview_only = preview_answer.lower().startswith("y")

    return GradingConfig(
        rubric_path=rubric_path,
        responses_path=responses_path,
        output_path=output_path,
        response_column=response_column,
        id_column=id_column,
        model=model,
        limit=limit,
        preview_only=preview_only,
        export_schema_path=export_schema_path or None,
        export_summary_path=export_summary_path or None,
    )


###################################################################################################
# ARGPARSE SUPPORT FOR VS CODE RUN CONFIGS
###################################################################################################

def parse_args() -> Optional[GradingConfig]:
    parser = argparse.ArgumentParser(description="AI Rubric Grader for VS Code")
    parser.add_argument("--rubric_path")
    parser.add_argument("--responses_path")
    parser.add_argument("--output_path")
    parser.add_argument("--response_column")
    parser.add_argument("--id_column")
    parser.add_argument("--rubric_sheet_name")
    parser.add_argument("--responses_sheet_name")
    parser.add_argument("--model", default=DEFAULT_MODEL)
    parser.add_argument("--checkpoint_interval", type=int, default=5)
    parser.add_argument("--limit", type=int)
    parser.add_argument("--preview_only", action="store_true")
    parser.add_argument("--export_schema_path")
    parser.add_argument("--export_summary_path")
    parser.add_argument("--interactive", action="store_true")

    args = parser.parse_args()

    # If user explicitly asked for interactive mode or omitted core args, use interactive setup
    if args.interactive or not (args.rubric_path and args.responses_path and args.output_path):
        return None

    return GradingConfig(
        rubric_path=args.rubric_path,
        responses_path=args.responses_path,
        output_path=args.output_path,
        response_column=args.response_column,
        id_column=args.id_column,
        rubric_sheet_name=args.rubric_sheet_name,
        responses_sheet_name=args.responses_sheet_name,
        model=args.model,
        checkpoint_interval=args.checkpoint_interval,
        limit=args.limit,
        preview_only=args.preview_only,
        export_schema_path=args.export_schema_path,
        export_summary_path=args.export_summary_path,
    )


###################################################################################################
# MAIN GRADING PIPELINE
###################################################################################################

def run_grading(config: GradingConfig) -> pd.DataFrame:
    rubric_path = resolve_path(config.rubric_path)
    responses_path = resolve_path(config.responses_path)
    output_path = resolve_path(config.output_path)

    schema_path = resolve_path(config.export_schema_path) if config.export_schema_path else output_path.with_name(output_path.stem + "_rubric_schema.json")
    summary_path = resolve_path(config.export_summary_path) if config.export_summary_path else output_path.with_name(output_path.stem + "_run_summary.json")
    checkpoint_path = output_path.with_name(output_path.stem + "_checkpoint.xlsx")

    print("\n" + "=" * 70)
    print("STARTING AI RUBRIC GRADER")
    print("=" * 70)

    client = get_client()

    print("\nLoading rubric...")
    rubric_text = read_rubric_file(rubric_path, sheet_name=config.rubric_sheet_name)
    print(f"Rubric text extracted: {len(rubric_text):,} characters")

    print("Parsing rubric schema with AI...")
    rubric = parse_rubric_with_ai(client, config.model, rubric_text, source_path=str(rubric_path))
    print(f"Rubric title: {rubric.title}")
    print(f"Criteria found: {len(rubric.criteria)}")
    print(f"Total max score: {rubric.get_max_total()}")
    for c in rubric.criteria:
        deps = f" | depends on: {', '.join(c.dependencies)}" if c.dependencies else ""
        print(f"  - {c.id}: {c.name} | max {c.max_score}{deps}")

    save_json(rubric_to_dict(rubric), schema_path)
    print(f"Rubric schema saved: {schema_path}")

    if config.preview_only:
        print("\nPreview only mode enabled. Exiting after rubric schema export.")
        save_json({
            "status": "preview_only",
            "timestamp": datetime.now().isoformat(),
            "rubric_path": str(rubric_path),
            "responses_path": str(responses_path),
            "output_path": str(output_path),
            "schema_path": str(schema_path),
        }, summary_path)
        return pd.DataFrame()

    system_prompt = build_dynamic_system_prompt(rubric)

    print("\nLoading responses...")
    df = read_responses_file(responses_path, sheet_name=config.responses_sheet_name)
    if config.limit:
        df = df.head(config.limit).copy()
    print(f"Rows to grade: {len(df)}")

    response_col = detect_response_column(df, config.response_column)
    id_col = detect_id_column(df, config.id_column)
    if id_col is None:
        id_col = "Response_ID"
        df[id_col] = range(1, len(df) + 1)

    print(f"Using response column: {response_col}")
    print(f"Using id column: {id_col}")

    df, col_map = prepare_output_columns(df, rubric)

    success = 0
    failed = 0
    review_count = 0
    skipped = 0
    start_time = time.time()

    for i in range(len(df)):
        response_id = df.at[i, id_col]
        response_text = str(df.at[i, response_col]).strip()

        if not response_text or response_text.lower() == "nan":
            skipped += 1
            df.at[i, "Overall_Rationale"] = "SKIPPED: empty response"
            df.at[i, "Needs_Review"] = "Empty response"
            logger.warning(f"Skipping empty response at row {i + 1}")
            continue

        try:
            print(f"Grading row {i + 1}/{len(df)} | ID {response_id}")
            result = grade_one_response(
                client=client,
                model=config.model,
                system_prompt=system_prompt,
                rubric=rubric,
                response_text=response_text,
                response_id=response_id,
            )

            for cid, (score_col, just_col) in col_map.items():
                df.at[i, score_col] = result["criteria_results"][cid]["score"]
                df.at[i, just_col] = result["criteria_results"][cid]["justification"]

            df.at[i, "Total_Score"] = result["total_score"]
            df.at[i, "Max_Possible_Score"] = result["max_possible_score"]
            df.at[i, "Overall_Rationale"] = result["overall_rationale"]

            review_flag = detect_borderline(result, rubric)
            df.at[i, "Needs_Review"] = review_flag
            if review_flag:
                review_count += 1
                print(f"  Review flag: {review_flag}")

            success += 1
            print(f"  Score: {result['total_score']}/{rubric.get_max_total()}")

        except Exception as e:
            failed += 1
            logger.error(f"Failed row {i + 1}, ID {response_id}: {e}")
            df.at[i, "Overall_Rationale"] = f"ERROR: {e}"
            df.at[i, "Needs_Review"] = "ERROR"
            print(f"  Error: {e}")

        if (i + 1) % config.checkpoint_interval == 0:
            safe_save_excel(df, checkpoint_path)
            print(f"Checkpoint saved: {checkpoint_path}")

        time.sleep(0.2)

    final_path = safe_save_excel(df, output_path)
    elapsed = round(time.time() - start_time, 2)

    summary = {
        "timestamp": datetime.now().isoformat(),
        "rubric_path": str(rubric_path),
        "responses_path": str(responses_path),
        "output_path": str(final_path),
        "schema_path": str(schema_path),
        "rows_total": len(df),
        "success": success,
        "failed": failed,
        "skipped": skipped,
        "review_count": review_count,
        "response_column": response_col,
        "id_column": id_col,
        "model": config.model,
        "elapsed_seconds": elapsed,
        "rubric_title": rubric.title,
        "rubric_total_max_score": rubric.get_max_total(),
    }
    save_json(summary, summary_path)

    print("\n" + "=" * 70)
    print("GRADING COMPLETE")
    print("=" * 70)
    print(f"Successful: {success}")
    print(f"Failed: {failed}")
    print(f"Skipped: {skipped}")
    print(f"Review flags: {review_count}")
    print(f"Time: {elapsed} seconds")
    print(f"Output saved: {final_path}")
    print(f"Run summary saved: {summary_path}")

    return df


###################################################################################################
# MAIN ENTRY
###################################################################################################

def main() -> None:
    print("\n" + "=" * 70)
    print("AI RUBRIC GRADER FOR VS CODE")
    print("=" * 70)
    print(f"Started: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    print(f"Workspace folder: {SCRIPT_DIR}")

    try:
        config = parse_args()
        if config is None:
            config = interactive_setup()

        print("\nConfiguration Summary")
        print(f"Rubric path: {resolve_path(config.rubric_path)}")
        print(f"Responses path: {resolve_path(config.responses_path)}")
        print(f"Output path: {resolve_path(config.output_path)}")
        print(f"Model: {config.model}")
        print(f"Limit: {config.limit if config.limit else 'All rows'}")
        print(f"Preview only: {config.preview_only}")

        if sys.stdin.isatty():
            confirm = input("\nProceed? (y/n): ").strip().lower()
            if confirm != "y":
                print("Cancelled.")
                return

        run_grading(config)

    except KeyboardInterrupt:
        print("\nInterrupted by user.")
        sys.exit(1)
    except Exception as e:
        logger.exception(f"Fatal error: {e}")
        print(f"\nFatal error: {e}")
        sys.exit(1)


if __name__ == "__main__":
    main()
